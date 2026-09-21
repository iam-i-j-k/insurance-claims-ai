from docx import Document
import os

def generate_policy():
    doc = Document()
    doc.add_heading('Comprehensive Motor Insurance Policy', 0)
    
    doc.add_heading('1. Coverage Limits', level=1)
    doc.add_paragraph('The maximum coverage limit for this motor policy is ₹1,50,000 per incident.')
    
    doc.add_heading('2. Deductibles', level=1)
    doc.add_paragraph('A mandatory deductible of ₹10,000 applies to all collision claims.')
    
    doc.add_heading('3. Filing Deadlines', level=1)
    doc.add_paragraph('All claims must be filed within 7 days of the incident occurrence. Claims filed after this period will be subject to manual review.')
    
    doc.add_heading('4. Mandatory Documentation', level=1)
    doc.add_paragraph('For any claim exceeding ₹1,00,000, an official First Information Report (FIR) from the police is strictly mandatory.')
    
    doc.add_heading('5. Fraud Risk Indicators', level=1)
    doc.add_paragraph('More than 2 claims in a 12-month period triggers an automatic fraud-risk review and potential Special Investigative Unit (SIU) escalation.')
    
    doc.save('../knowledge_base/motor/motor_policy.docx')

def generate_claim_form():
    doc = Document()
    doc.add_heading('Motor Claim Form - Submission', 0)
    
    doc.add_paragraph('Policyholder Name: Rajesh Kumar')
    doc.add_paragraph('Policy Number: MOT-2026-9912')
    doc.add_paragraph('Claim Type: Motor / Collision')
    doc.add_paragraph('Incident Date: 10/09/2026')
    doc.add_paragraph('Filing Date: 15/09/2026 (Filed 5 days after incident)')
    doc.add_paragraph('Incident Location: Main Highway, Mumbai')
    
    doc.add_paragraph('Description of Incident: Rear-ended by an unknown vehicle while stopped at a traffic light. The bumper and trunk are severely damaged.')
    
    doc.add_paragraph('Claimed Amount: ₹1,80,000')
    
    doc.add_paragraph('Supporting Documents Submitted: Repair Estimate from AutoCare Garage, Photographs of the vehicle damage.')
    doc.add_paragraph('FIR Attached: No')
    
    doc.add_paragraph('Prior Claims: I have had 3 prior claims in the last 12 months for minor scratches.')
    
    doc.save('demo_claim_form.docx')

if __name__ == "__main__":
    generate_policy()
    generate_claim_form()
    print("Demo documents generated successfully.")
