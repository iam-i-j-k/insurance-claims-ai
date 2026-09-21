from docx import Document

def create_claim_doc():
    doc = Document()
    
    doc.add_heading('First Information Report (FIR) & Medical Summary', 0)
    
    doc.add_heading('Claimant Details', level=1)
    doc.add_paragraph('Name: Michael Scott')
    doc.add_paragraph('Policy Number: POL-H-987654')
    doc.add_paragraph('Date of Incident: October 20, 2026')
    
    doc.add_heading('Incident Description', level=1)
    doc.add_paragraph(
        'On the evening of October 20, the claimant experienced severe abdominal pain and was rushed to '
        'Scranton General Hospital. After initial tests, it was determined that the claimant suffered from '
        'acute appendicitis requiring immediate surgical intervention.'
    )
    
    doc.add_heading('Medical Diagnosis & Treatment', level=1)
    doc.add_paragraph('Diagnosis: Acute Appendicitis')
    doc.add_paragraph('Procedure: Emergency Appendectomy (Laparoscopic)')
    doc.add_paragraph('Hospitalization Period: Oct 20, 2026 - Oct 23, 2026')
    
    doc.add_heading('Financial Summary (Claimed Amount)', level=1)
    doc.add_paragraph('Surgeon Fees: $4,500.00')
    doc.add_paragraph('Anesthesia: $1,200.00')
    doc.add_paragraph('Room & Board (3 days): $3,000.00')
    doc.add_paragraph('Total Claimed Amount: $8,700.00')
    
    doc.add_heading('Attending Physician Statement', level=1)
    doc.add_paragraph(
        'I certify that Michael Scott was treated for acute appendicitis. This is a sudden, unforeseen medical '
        'emergency and not a pre-existing condition. The procedure was medically necessary.'
    )
    doc.add_paragraph('Sign: Dr. Jim Halpert, MD')
    
    doc.save('Michael_Scott_Health_Claim.docx')
    print("Created Michael_Scott_Health_Claim.docx")

if __name__ == "__main__":
    create_claim_doc()
